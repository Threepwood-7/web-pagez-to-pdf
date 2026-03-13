"""Export writers for PDF/images/TIFF/DOCX/PPTX/XLSX from capture session data."""

from __future__ import annotations

import logging
import re
from contextlib import suppress
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, ClassVar, Protocol, cast

from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

from .image_processing import (
    DEFAULT_CONTENT_SIZING_MODE,
    PAPER_SIZES,
    PageSlice,
    apply_edit_transform,
    compute_page_slices,
    content_points_per_pixel,
    normalize_content_sizing_mode,
)
from .models import EditAdjustments

if TYPE_CHECKING:
    from .models import CaptureItem, ExportRequest

try:
    from docx import Document
    from docx.shared import Mm
except ImportError:  # pragma: no cover
    Document = None  # type: ignore[assignment]
    Mm = None  # type: ignore[assignment]

try:
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:  # pragma: no cover
    Presentation = None  # type: ignore[assignment]
    Inches = None  # type: ignore[assignment]

try:
    import xlsxwriter
except ImportError:  # pragma: no cover
    xlsxwriter = None  # type: ignore[assignment]

EXPORT_LOGGER = logging.getLogger("web_pagez_to_pdf.export")


class _RichTextProbe(HTMLParser):
    """Detect whether an HTML fragment contains meaningful visible content."""

    _VOID_MEDIA_TAGS: ClassVar[set[str]] = {
        "img",
        "hr",
        "svg",
        "canvas",
        "video",
        "audio",
        "object",
        "iframe",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._suppress_depth = 0
        self.has_visible_text = False
        self.has_media_content = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        name = str(tag or "").strip().lower()
        if name in {"style", "script", "head"}:
            self._suppress_depth += 1
            return
        if name in self._VOID_MEDIA_TAGS:
            self.has_media_content = True

    def handle_endtag(self, tag: str) -> None:
        name = str(tag or "").strip().lower()
        if name in {"style", "script", "head"} and self._suppress_depth > 0:
            self._suppress_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._suppress_depth > 0:
            return
        if re.sub(r"\s+", "", str(data or "")):
            self.has_visible_text = True

    def has_meaningful_content(self) -> bool:
        return self.has_visible_text or self.has_media_content


@dataclass(slots=True)
class PageFrame:
    """Renderable page frame with source metadata."""

    image: Image.Image
    title: str
    source_item: CaptureItem
    page_index_in_item: int
    page_count_in_item: int
    content_points_per_pixel: float = 0.0


@dataclass(slots=True)
class ExportResult:
    """Paths written by one export invocation."""

    generated_paths: list[Path]


@dataclass(slots=True)
class _XlsxExportRow:
    title: str
    item_id: str
    source_image_path: str
    page_index: int
    page_count: int
    image: Image.Image


class _PdfCanvas(Protocol):
    """Subset of reportlab canvas methods used by this exporter module."""

    def drawInlineImage(
        self,
        image: Image.Image,
        x: float,
        y: float,
        width: float | None = None,
        height: float | None = None,
    ) -> object: ...


def run_export(request: ExportRequest) -> ExportResult:
    """Export selected capture queue into requested output formats."""

    requested_formats = [
        name
        for name, enabled in (
            ("pdf", request.formats.pdf),
            ("paged_images", request.formats.paged_images),
            ("long_image", request.formats.long_image),
            ("tiff", request.formats.tiff),
            ("docx", request.formats.docx),
            ("pptx", request.formats.pptx),
            ("xlsx", request.formats.xlsx),
        )
        if enabled
    ]
    EXPORT_LOGGER.info(
        "run_export start captures=%s combine_mode=%s formats=%s "
        "output_dir=%s basename=%s",
        len(request.captures),
        request.combine_mode,
        ",".join(requested_formats) or "none",
        request.output_dir,
        request.basename,
    )
    request.output_dir = Path(request.output_dir)
    request.output_dir.mkdir(parents=True, exist_ok=True)
    frames, transformed_images = build_page_frames(request)
    if not frames:
        raise RuntimeError("No exportable frames are available.")

    generated: list[Path] = []
    if request.formats.pdf:
        EXPORT_LOGGER.info("writer start format=pdf")
        generated.append(export_pdf(request, frames))
        EXPORT_LOGGER.info("writer done format=pdf path=%s", generated[-1])
    if request.formats.paged_images:
        EXPORT_LOGGER.info("writer start format=paged_images")
        generated.extend(export_paged_images(request, frames))
        EXPORT_LOGGER.info("writer done format=paged_images")
    if request.formats.long_image:
        EXPORT_LOGGER.info("writer start format=long_image")
        generated.append(export_long_image(request, transformed_images))
        EXPORT_LOGGER.info("writer done format=long_image path=%s", generated[-1])
    if request.formats.tiff:
        EXPORT_LOGGER.info("writer start format=tiff")
        generated.append(export_multipage_tiff(request, frames))
        EXPORT_LOGGER.info("writer done format=tiff path=%s", generated[-1])
    if request.formats.docx:
        EXPORT_LOGGER.info("writer start format=docx")
        generated.append(export_docx(request, frames, transformed_images))
        EXPORT_LOGGER.info("writer done format=docx path=%s", generated[-1])
    if request.formats.pptx:
        EXPORT_LOGGER.info("writer start format=pptx")
        generated.append(export_pptx(request, frames, transformed_images))
        EXPORT_LOGGER.info("writer done format=pptx path=%s", generated[-1])
    if request.formats.xlsx:
        EXPORT_LOGGER.info("writer start format=xlsx")
        generated.append(export_xlsx(request, frames, transformed_images))
        EXPORT_LOGGER.info("writer done format=xlsx path=%s", generated[-1])
    EXPORT_LOGGER.info(
        "run_export done generated=%s paths=%s",
        len(generated),
        [str(path) for path in generated],
    )
    return ExportResult(generated_paths=generated)


def build_page_frames(
    request: ExportRequest,
) -> tuple[list[PageFrame], list[tuple[CaptureItem, Image.Image]]]:
    """Build transformed images and split frames according to export mode."""

    transformed_images: list[tuple[CaptureItem, Image.Image]] = []
    for item in request.captures:
        image = Image.open(item.image_path).convert("RGB")
        edits = _edits_for_item(request, item)
        transformed = apply_edit_transform(
            image, request.layout, edits, include_scale=False
        )
        transformed_images.append((item, transformed))

    frames: list[PageFrame] = []
    if request.combine_mode:
        for item, image in transformed_images:
            edits = _edits_for_item(request, item)
            content_mode = _content_sizing_mode_for_edits(edits)
            scale_percent = _scale_percent_for_edits(edits)
            points_per_px = content_points_per_pixel(
                image.width,
                image.height,
                request.layout,
                content_mode,
                scale_percent=scale_percent,
            )
            slices = compute_page_slices(
                image,
                request.layout,
                edits.split_markers_px,
                content_sizing_mode=content_mode,
                scale_percent=scale_percent,
            )
            frames.extend(
                _frames_for_slices(
                    item,
                    image,
                    slices,
                    content_points_per_pixel=points_per_px,
                )
            )
    else:
        first_item, first_image = transformed_images[0]
        edits = _edits_for_item(request, first_item)
        content_mode = _content_sizing_mode_for_edits(edits)
        scale_percent = _scale_percent_for_edits(edits)
        points_per_px = content_points_per_pixel(
            first_image.width,
            first_image.height,
            request.layout,
            content_mode,
            scale_percent=scale_percent,
        )
        slices = compute_page_slices(
            first_image,
            request.layout,
            edits.split_markers_px,
            content_sizing_mode=content_mode,
            scale_percent=scale_percent,
        )
        frames.extend(
            _frames_for_slices(
                first_item,
                first_image,
                slices,
                content_points_per_pixel=points_per_px,
            )
        )
    return (frames, transformed_images)


def _edits_for_item(request: ExportRequest, item: CaptureItem) -> EditAdjustments:
    return (
        request.edits_by_item_id.get(item.item_id) or request.edits or EditAdjustments()
    )


def _content_sizing_mode_for_edits(edits: EditAdjustments) -> str:
    mode_op = edits.get_operation("content_sizing_mode")
    mode_value = (
        mode_op.params.get("mode")
        if mode_op is not None
        else DEFAULT_CONTENT_SIZING_MODE
    )
    return normalize_content_sizing_mode(mode_value)


def _scale_percent_for_edits(edits: EditAdjustments) -> float:
    scale_op = edits.get_operation("scale")
    if scale_op is None:
        return 100.0
    try:
        return max(10.0, float(scale_op.params.get("percent", 100.0)))
    except (TypeError, ValueError):
        return 100.0


def _frames_for_slices(
    item: CaptureItem,
    image: Image.Image,
    slices: list[PageSlice],
    *,
    content_points_per_pixel: float,
) -> list[PageFrame]:
    frames: list[PageFrame] = []
    page_total = len(slices)
    for idx, slice_info in enumerate(slices):
        frame = image.crop((0, slice_info.top, image.width, slice_info.bottom))
        frames.append(
            PageFrame(
                image=frame,
                title=item.title,
                source_item=item,
                page_index_in_item=idx + 1,
                page_count_in_item=page_total,
                content_points_per_pixel=float(content_points_per_pixel),
            )
        )
    return frames


def export_pdf(request: ExportRequest, frames: list[PageFrame]) -> Path:
    """Write PDF with optional rich header/footer."""

    page_size = PAPER_SIZES.get(request.layout.paper_name.upper(), A4)
    page_w, page_h = page_size
    if request.layout.orientation.lower() == "landscape":
        page_w, page_h = page_h, page_w

    output_path = request.output_dir / f"{request.basename}.pdf"
    pdf = canvas.Canvas(str(output_path), pagesize=(page_w, page_h))
    pdf_writer = cast("_PdfCanvas", pdf)
    total_pages = len(frames)
    avail_w = (
        page_w
        - (
            request.layout.margin_left_mm
            + request.layout.margin_right_mm
            + request.layout.gutter_mm
        )
        * mm
    )
    content_x = (request.layout.margin_left_mm + request.layout.gutter_mm) * mm
    for page_index, frame in enumerate(frames, start=1):
        points_per_px = float(frame.content_points_per_pixel)
        if points_per_px <= 0.0:
            points_per_px = float(avail_w) / float(max(1, int(frame.image.width)))
        rendered_w = float(frame.image.width) * points_per_px
        rendered_h = float(frame.image.height) * points_per_px
        y = page_h - request.layout.margin_top_mm * mm - rendered_h
        pdf_writer.drawInlineImage(
            frame.image, content_x, y, width=rendered_w, height=rendered_h
        )

        context = {
            "title": frame.title,
            "datetime": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "page": str(page_index),
            "pages": str(total_pages),
        }
        _draw_rich_text(
            pdf,
            request.layout.header_rich_text,
            context,
            content_x,
            page_h - 16,
            avail_w,
        )
        _draw_rich_text(
            pdf, request.layout.footer_rich_text, context, content_x, 12, avail_w
        )
        if page_index < total_pages:
            pdf.showPage()
            pdf.setPageSize((page_w, page_h))
    pdf.save()
    return output_path


def _draw_rich_text(
    pdf: canvas.Canvas,
    rich_text: str,
    context: dict[str, str],
    x_pos: float,
    y_pos: float,
    width: float,
) -> None:
    source = _meaningful_rich_text_or_empty(rich_text)
    text = _apply_tokens(source, context).strip()
    if not text:
        return
    paragraph = Paragraph(text)
    wrapped_w, wrapped_h = paragraph.wrap(width, 60)
    paragraph.drawOn(pdf, x_pos, y_pos - wrapped_h)
    if wrapped_w <= 0:
        return


def export_paged_images(request: ExportRequest, frames: list[PageFrame]) -> list[Path]:
    """Export each page frame as numbered PNG files."""

    generated: list[Path] = []
    for index, frame in enumerate(frames, start=1):
        filename = f"{request.basename}_p{index:03d}.png"
        path = request.output_dir / filename
        frame.image.save(path, format="PNG")
        generated.append(path)
    return generated


def export_long_image(
    request: ExportRequest, transformed: list[tuple[CaptureItem, Image.Image]]
) -> Path:
    """Export one long stitched image for selected captures."""

    if request.combine_mode:
        width = max(img.width for _item, img in transformed)
        total_height = (
            sum(img.height for _item, img in transformed)
            + max(0, len(transformed) - 1) * 16
        )
        canvas_img = Image.new("RGB", (width, total_height), "white")
        y_pos = 0
        for _item, image in transformed:
            canvas_img.paste(image, (0, y_pos))
            y_pos += image.height + 16
    else:
        canvas_img = transformed[0][1]
    output = request.output_dir / f"{request.basename}_long.png"
    canvas_img.save(output, format="PNG")
    return output


def export_multipage_tiff(request: ExportRequest, frames: list[PageFrame]) -> Path:
    """Export split pages into one multi-page TIFF file."""

    output = request.output_dir / f"{request.basename}.tiff"
    if not frames:
        raise RuntimeError("No frames available for TIFF export.")
    head = frames[0].image.copy().convert("RGB")
    tail = [frame.image.copy().convert("RGB") for frame in frames[1:]]
    try:
        head.save(
            output,
            format="TIFF",
            save_all=True,
            append_images=tail,
            compression="tiff_deflate",
        )
        return output
    except TypeError:
        # Pillow/libtiff can fail after prior PNG writes in the same run.
        with suppress(OSError):
            output.unlink()
        fallback_head = frames[0].image.copy().convert("RGB")
        fallback_tail = [frame.image.copy().convert("RGB") for frame in frames[1:]]
        try:
            fallback_head.save(
                output,
                format="TIFF",
                save_all=True,
                append_images=fallback_tail,
            )
            return output
        except Exception as exc:  # pragma: no cover - depends on local Pillow/libtiff
            raise RuntimeError(f"TIFF export failed after fallback: {exc}") from exc
    except Exception as exc:
        raise RuntimeError(f"TIFF export failed: {exc}") from exc


def export_docx(
    request: ExportRequest,
    frames: list[PageFrame],
    transformed: list[tuple[CaptureItem, Image.Image]],
) -> Path:
    """Export images into DOCX with user-selected layout mode."""

    if Document is None or Mm is None:  # pragma: no cover
        raise RuntimeError("python-docx dependency is required for DOCX export.")
    document = Document()
    output = request.output_dir / f"{request.basename}.docx"
    mode = request.docx_pptx_mode.strip().lower()
    if mode == "per_capture":
        sources = [image for _item, image in transformed]
    else:
        sources = [frame.image for frame in frames]
    for index, image in enumerate(sources):
        if index > 0:
            document.add_page_break()
        paragraph = document.add_paragraph()
        with BytesIO() as buffer:
            image.save(buffer, format="PNG")
            buffer.seek(0)
            paragraph.add_run().add_picture(buffer, width=Mm(170))
    document.save(str(output))
    return output


def export_pptx(
    request: ExportRequest,
    frames: list[PageFrame],
    transformed: list[tuple[CaptureItem, Image.Image]],
) -> Path:
    """Export images into PPTX slides with selected layout mode."""

    if Presentation is None or Inches is None:  # pragma: no cover
        raise RuntimeError("python-pptx dependency is required for PPTX export.")
    presentation = Presentation()
    output = request.output_dir / f"{request.basename}.pptx"
    mode = request.docx_pptx_mode.strip().lower()
    if mode == "per_capture":
        sources = [image for _item, image in transformed]
    else:
        sources = [frame.image for frame in frames]
    for image in sources:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        with BytesIO() as buffer:
            image.save(buffer, format="PNG")
            buffer.seek(0)
            slide.shapes.add_picture(
                buffer, Inches(0.3), Inches(0.3), width=Inches(12.7)
            )
    presentation.save(str(output))
    return output


def export_xlsx(
    request: ExportRequest,
    frames: list[PageFrame],
    transformed: list[tuple[CaptureItem, Image.Image]],
) -> Path:
    """Export split pages/captures into one XLSX sheet with metadata and previews."""

    if xlsxwriter is None:  # pragma: no cover
        raise RuntimeError("xlsxwriter dependency is required for XLSX export.")
    output = request.output_dir / f"{request.basename}.xlsx"
    mode = request.docx_pptx_mode.strip().lower()
    rows = _xlsx_export_rows(mode, frames, transformed)
    workbook = xlsxwriter.Workbook(str(output))
    image_streams: list[BytesIO] = []
    try:
        worksheet = workbook.add_worksheet("captures")
        header_format = workbook.add_format({"bold": True})
        body_format = workbook.add_format({"valign": "top"})
        headers = [
            "title",
            "item_id",
            "source_image_path",
            "page_index",
            "page_count",
            "image_pixel_size",
            "preview_image",
        ]
        for column, header in enumerate(headers):
            worksheet.write(0, column, header, header_format)
        worksheet.freeze_panes(1, 0)
        worksheet.set_column(0, 0, 28)
        worksheet.set_column(1, 1, 20)
        worksheet.set_column(2, 2, 56)
        worksheet.set_column(3, 4, 11)
        worksheet.set_column(5, 5, 16)
        worksheet.set_column(6, 6, 46)

        max_preview_width_px = 320
        for row_index, row in enumerate(rows, start=1):
            worksheet.write(row_index, 0, row.title, body_format)
            worksheet.write(row_index, 1, row.item_id, body_format)
            worksheet.write(row_index, 2, row.source_image_path, body_format)
            worksheet.write_number(row_index, 3, float(row.page_index), body_format)
            worksheet.write_number(row_index, 4, float(row.page_count), body_format)
            worksheet.write(
                row_index,
                5,
                f"{row.image.width}x{row.image.height}",
                body_format,
            )
            scale = 1.0
            if row.image.width > max_preview_width_px:
                scale = float(max_preview_width_px) / float(row.image.width)
            preview_height_px = max(1.0, float(row.image.height) * scale)
            worksheet.set_row(row_index, max(20.0, preview_height_px * 0.75 + 4.0))
            stream = BytesIO()
            row.image.save(stream, format="PNG")
            stream.seek(0)
            image_streams.append(stream)
            worksheet.insert_image(
                row_index,
                6,
                "preview.png",
                {
                    "image_data": stream,
                    "x_scale": scale,
                    "y_scale": scale,
                    "x_offset": 2,
                    "y_offset": 2,
                },
            )
    finally:
        workbook.close()
    return output


def _xlsx_export_rows(
    mode: str,
    frames: list[PageFrame],
    transformed: list[tuple[CaptureItem, Image.Image]],
) -> list[_XlsxExportRow]:
    if mode == "per_capture":
        return [
            _XlsxExportRow(
                title=item.title,
                item_id=item.item_id,
                source_image_path=str(item.image_path),
                page_index=1,
                page_count=1,
                image=image,
            )
            for item, image in transformed
        ]
    return [
        _XlsxExportRow(
            title=frame.title,
            item_id=frame.source_item.item_id,
            source_image_path=str(frame.source_item.image_path),
            page_index=frame.page_index_in_item,
            page_count=frame.page_count_in_item,
            image=frame.image,
        )
        for frame in frames
    ]


def sanitize_basename(value: str) -> str:
    """Filesystem-safe export base stem."""

    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip())
    cleaned = cleaned.strip("-._")
    if not cleaned:
        return "capture"
    return cleaned[:140]


def _apply_tokens(text: str, context: dict[str, str]) -> str:
    result = text
    for key, value in context.items():
        result = result.replace(f"{{{key}}}", value)
    return result


def _meaningful_rich_text_or_empty(rich_text: str) -> str:
    source = str(rich_text or "")
    if not source.strip():
        return ""
    probe = _RichTextProbe()
    with suppress(Exception):
        probe.feed(source)
        probe.close()
    if probe.has_meaningful_content():
        return source
    return ""
